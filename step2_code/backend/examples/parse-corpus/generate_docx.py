"""Generate sample.docx for parse-corpus examples.

Run from the parse-corpus directory:
    cd examples/parse-corpus
    python generate_docx.py

Structure is reproducible (same paragraphs, tables, text in same order).
ZIP byte-level reproducibility is NOT guaranteed — python-docx/lxml may
emit different attribute ordering or timestamps across versions.
After running, update the sha256 in manifest.json:
    python -c "import hashlib; print(hashlib.sha256(open('sample.docx','rb').read()).hexdigest())"
"""

from pathlib import Path

from docx import Document

OUT_DIR = Path(__file__).parent
OUT_PATH = OUT_DIR / "sample.docx"

doc = Document()

# Paragraph 1
doc.add_paragraph("【合成测试文件 — 不得用于论文评测或真实校园问答】")

# Paragraph 2
doc.add_paragraph("四川农业大学合成测试课程成绩登记表（样例）")

# Table 1: grade records
table = doc.add_table(rows=4, cols=4)
table.style = "Table Grid"
headers = ["学号", "姓名", "课程", "成绩"]
for i, h in enumerate(headers):
    table.rows[0].cells[i].text = h
data = [
    ["20240001", "张三", "高等数学A1", "85"],
    ["20240002", "李四", "高等数学A1", "92"],
    ["20240003", "王五", "高等数学A1", "78"],
]
for r, row_data in enumerate(data):
    for c, val in enumerate(row_data):
        table.rows[r + 1].cells[c].text = val

# Empty paragraph (skipped in output but index counts)
doc.add_paragraph("")

# Paragraph 3
doc.add_paragraph(
    "以上成绩仅供参考。实际成绩以教务系统公布为准。"
    "如有疑问，请在成绩公布后10个工作日内向所在学院提出复核申请。"
)

# Table 2: requirements
table2 = doc.add_table(rows=3, cols=3)
table2.style = "Table Grid"
headers2 = ["项目", "截止日期", "备注"]
for i, h in enumerate(headers2):
    table2.rows[0].cells[i].text = h
reqs = [
    ["缓考申请", "考试前3个工作日", "需医院证明"],
    ["成绩复核", "公布后10个工作日", "仅查分不查卷"],
]
for r, row_data in enumerate(reqs):
    for c, val in enumerate(row_data):
        table2.rows[r + 1].cells[c].text = val

# Paragraph 4
doc.add_paragraph("本文件为合成测试数据，不代表真实教务政策。")

doc.save(str(OUT_PATH))
print(f"Created: {OUT_PATH}")
