"""Offline document body parsing.

Converts validated TXT/DOCX files into body blocks with source locators.
Reuses T003 validation; does not generate answers, vectors, or chunks.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import sys
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from docx.oxml.ns import qn

from campus_assistant.corpus.validate import (
    MANIFEST_READ_ERROR,
    validate_manifest,
)

# Error codes for parsing
DOCUMENT_PARSE_ERROR = "DOCUMENT_PARSE_ERROR"
EMPTY_DOCUMENT = "EMPTY_DOCUMENT"
UNSUPPORTED_FORMAT = "UNSUPPORTED_FORMAT"

PARSER_VERSION = "1.0"
# PDF 独立解析器版本：只参与 PDF 块 ID 计算。全局 PARSER_VERSION 保持 1.0，
# TXT/DOCX 块 ID 与既有模板/规则引用不受影响（T021 红线）。
PDF_PARSER_VERSION = "pdf-1.0"
SUPPORTED_EXTENSIONS = {".txt", ".docx", ".pdf"}

_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _error(code: str, document_id: str | None, field: str, message: str) -> dict:
    return {"code": code, "document_id": document_id, "field": field, "message": message}


def _make_block_id(
    document_id: str, sha256: str, kind: str, locator_key: str,
    parser_version: str = PARSER_VERSION,
) -> str:
    """Deterministic block ID from document identity and block location.

    parser_version participates in the seed; TXT/DOCX keep the default
    (PARSER_VERSION) so existing block IDs never change. PDF blocks pass
    PDF_PARSER_VERSION (T021).
    """
    seed = f"{document_id}|{sha256}|{parser_version}|{kind}|{locator_key}"
    return hashlib.sha256(seed.encode("utf-8")).hexdigest()[:16]


# ---------------------------------------------------------------------------
# TXT parser
# ---------------------------------------------------------------------------

def parse_txt(path: Path, document_id: str, sha256: str) -> tuple[list[dict], list[dict]]:
    """Parse a UTF-8 TXT file into paragraph blocks split by blank lines.

    Returns (blocks, errors). Empty blocks list with an error = failure.
    """
    try:
        raw = path.read_bytes()
    except PermissionError:
        return [], [_error(DOCUMENT_PARSE_ERROR, document_id, "local_path", "无读取权限")]
    except OSError:
        return [], [_error(DOCUMENT_PARSE_ERROR, document_id, "local_path", "读取文件失败")]

    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return [], [_error(DOCUMENT_PARSE_ERROR, document_id, "local_path", "文件编码不是 UTF-8")]

    text = text.replace("\r\n", "\n").replace("\r", "\n")

    if not text.strip():
        return [], [_error(EMPTY_DOCUMENT, document_id, "local_path", "文件内容为空")]

    lines = text.split("\n")

    paragraphs: list[tuple[int, int, str]] = []
    current_lines: list[str] = []
    current_start: int | None = None

    for i, line in enumerate(lines):
        line_num = i + 1
        if line.strip() == "":
            if current_lines:
                paragraphs.append((current_start, line_num - 1, "\n".join(current_lines)))
                current_lines = []
                current_start = None
        else:
            if current_start is None:
                current_start = line_num
            current_lines.append(line)

    if current_lines:
        paragraphs.append((current_start, len(lines), "\n".join(current_lines)))

    blocks: list[dict] = []
    for ordinal, (line_start, line_end, para_text) in enumerate(paragraphs, start=1):
        if not para_text.strip():
            continue
        locator_key = f"lines:{line_start}-{line_end}"
        blocks.append({
            "block_id": _make_block_id(document_id, sha256, "paragraph", locator_key),
            "ordinal": ordinal,
            "kind": "paragraph",
            "text": para_text,
            "locator": {"type": "lines", "line_start": line_start, "line_end": line_end},
        })

    if not blocks:
        return [], [_error(EMPTY_DOCUMENT, document_id, "local_path", "文件无有效正文段落")]

    return blocks, []


# ---------------------------------------------------------------------------
# DOCX parser
# ---------------------------------------------------------------------------

def _cell_text(cell: Any) -> str:
    return "\n".join(p.text for p in cell.paragraphs).strip()


def _detect_merged_cells(table_obj: Any) -> list[str]:
    """Detect horizontally/vertically merged cells in a table.

    Merge attributes are inside w:tcPr (cell properties), not directly under w:tc.
    Returns list of warning strings.
    """
    warnings: list[str] = []
    tbl = table_obj._tbl
    for row_idx, tr in enumerate(tbl.findall(qn("w:tr")), start=1):
        for cell_idx, tc in enumerate(tr.findall(qn("w:tc")), start=1):
            tc_pr = tc.find(qn("w:tcPr"))
            if tc_pr is None:
                continue
            if tc_pr.find(qn("w:gridSpan")) is not None:
                warnings.append(f"表格包含水平合并单元格（第 {row_idx} 行第 {cell_idx} 列），文本可能不完整")
            vmerge = tc_pr.find(qn("w:vMerge"))
            if vmerge is not None:
                val = vmerge.get(qn("w:val"), "continue")
                if val != "continue":
                    warnings.append(f"表格包含垂直合并单元格（第 {row_idx} 行第 {cell_idx} 列），文本可能不完整")
    return warnings


def _detect_header_footer(doc: Any) -> list[str]:
    """Detect non-empty headers and footers. Returns warning list."""
    warnings: list[str] = []
    for section in doc.sections:
        for hf, label in [
            (section.header, "页眉"),
            (section.footer, "页脚"),
            (section.even_page_header, "偶数页眉"),
            (section.even_page_footer, "偶数页脚"),
            (section.first_page_header, "首页页眉"),
            (section.first_page_footer, "首页页脚"),
        ]:
            if hf is not None:
                text = "\n".join(p.text for p in hf.paragraphs).strip()
                if text:
                    warnings.append(f"文档包含非空{label}，本解析器不提取{label}内容")
    return warnings


def _detect_floating_images(doc: Any) -> list[str]:
    """Detect floating/drawing images (not inline). Returns warning list."""
    warnings: list[str] = []
    body_xml = doc.element.body
    # Look for w:drawing/wp:anchor (floating) vs wp:inline
    for drawing in body_xml.iter(qn("w:drawing")):
        for anchor in drawing.iter(qn("wp:anchor")):
            warnings.append("文档包含浮动图片，本解析器不提取图片内容")
            return warnings  # report once
    return warnings


def parse_docx(
    path: Path, document_id: str, sha256: str
) -> tuple[list[dict], list[dict], list[str]]:
    """Parse a DOCX file preserving paragraph/table interleaving.

    Returns (blocks, errors, warnings).
    """
    warnings: list[str] = []

    try:
        doc = DocxDocument(str(path))
    except Exception:
        return [], [
            _error(DOCUMENT_PARSE_ERROR, document_id, "local_path", "DOCX 文件损坏或格式不正确")
        ], []

    # R2: Wrap entire iteration to catch parse-time exceptions
    try:
        blocks: list[dict] = []
        ordinal = 0
        paragraph_index = 0
        table_index = 0

        body = doc.element.body
        for child in body:
            tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag

            if tag == "p":
                paragraph_index += 1
                para_text = None
                for p in doc.paragraphs:
                    if p._element is child:
                        para_text = p.text.strip()
                        break
                if para_text is None:
                    para_text = "".join(child.itertext()).strip()

                if not para_text:
                    continue

                ordinal += 1
                locator_key = f"para:{paragraph_index}"
                blocks.append({
                    "block_id": _make_block_id(document_id, sha256, "paragraph", locator_key),
                    "ordinal": ordinal,
                    "kind": "paragraph",
                    "text": para_text,
                    "locator": {"type": "paragraph", "paragraph_index": paragraph_index},
                })

            elif tag == "tbl":
                table_index += 1
                table_obj = None
                for t in doc.tables:
                    if t._tbl is child:
                        table_obj = t
                        break
                if table_obj is None:
                    warnings.append(f"表格 {table_index} 无法解析")
                    continue

                # R4: Detect merged cells
                merge_warnings = _detect_merged_cells(table_obj)
                warnings.extend(merge_warnings)

                # R1: row_index counts ALL rows (including empty), ordinal only for output
                source_row_index = 0
                for row in table_obj.rows:
                    source_row_index += 1
                    cells_text = [_cell_text(cell) for cell in row.cells]
                    row_text = "\t".join(cells_text)

                    if not row_text.strip():
                        continue  # skip empty rows in output

                    ordinal += 1
                    locator_key = f"table:{table_index},row:{source_row_index}"
                    blocks.append({
                        "block_id": _make_block_id(
                            document_id, sha256, "table_row", locator_key
                        ),
                        "ordinal": ordinal,
                        "kind": "table_row",
                        "text": row_text,
                        "locator": {
                            "type": "table_row",
                            "table_index": table_index,
                            "row_index": source_row_index,
                        },
                        "cells": cells_text,
                    })

                # R4: Detect nested tables
                if any(cell.tables for row in table_obj.rows for cell in row.cells):
                    warnings.append(f"表格 {table_index} 包含嵌套表格，可能未完整提取")

    except Exception as e:
        # R1: Isolate per-document parse exceptions including third-party XML errors
        # (e.g. docx.oxml.exceptions.InvalidXmlError from malformed gridSpan)
        return [], [
            _error(DOCUMENT_PARSE_ERROR, document_id, "local_path",
                   f"DOCX 解析异常: {type(e).__name__}")
        ], warnings

    # R4: Detect headers, footers, floating images, inline images
    try:
        warnings.extend(_detect_header_footer(doc))
        warnings.extend(_detect_floating_images(doc))
        if doc.inline_shapes:
            warnings.append("文档包含内联图片，本解析器不提取图片内容")
    except Exception:
        warnings.append("检测页眉页脚或图片时发生异常，部分内容可能未提示")

    # R4: Always include warnings key (empty list if no warnings)
    if not blocks:
        return [], [_error(EMPTY_DOCUMENT, document_id, "local_path", "文档无有效正文内容")], warnings

    return blocks, [], warnings


# ---------------------------------------------------------------------------
# Main parse logic
# ---------------------------------------------------------------------------

def _resolve_path(corpus_root: Path, local_path: str) -> Path | None:
    try:
        resolved_root = corpus_root.resolve()
        target = (resolved_root / local_path).resolve()
        target.relative_to(resolved_root)
        return target
    except (ValueError, OSError):
        return None


# ---------------------------------------------------------------------------
# PDF parser (T021: text PDFs only — pypdf, no OCR)
# ---------------------------------------------------------------------------

def _normalize_pdf_text(text: str) -> str:
    """Deterministic whitespace normalization for extracted PDF text.

    规则（可复查）：逐行去除尾部空白；换行统一为 \\n；连续空行折叠为一个
    空行（保留段落分隔）；不改写任何非空白字符，不补正文。
    """
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    normalized = "\n".join(lines)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)
    return normalized.strip("\n")


def parse_pdf(path: Path, document_id: str, sha256: str) -> tuple[list[dict], list[dict], list[str]]:
    """Extract text blocks from a text PDF, one group per page part.

    Block locator: {"type": "pdf_page", "page": N, "part": K}（1 基页码 +
    页内段序，页内稳定定位）。提取文本即块原文（规范化后），引用跨度相对
    该提取块逐字一致；不宣称与视觉排版等价；不使用 OCR 或 LLM 补正文。
    """
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover
        return [], [_error(DOCUMENT_PARSE_ERROR, document_id, "local_path", "缺少 PDF 解析组件（pypdf）")], []

    try:
        reader = PdfReader(str(path))
        if getattr(reader, "is_encrypted", False):
            return [], [_error(
                DOCUMENT_PARSE_ERROR, document_id, "local_path",
                "PDF 已加密，不提供解密绕过；请提供未加密的文本 PDF",
            )], []
    except Exception as e:
        return [], [_error(
            DOCUMENT_PARSE_ERROR, document_id, "local_path",
            f"PDF 无法打开（可能损坏或加密）：{type(e).__name__}",
        )], []

    blocks: list[dict] = []
    pages_without_text: list[int] = []
    ordinal = 0
    for page_index, page in enumerate(reader.pages, start=1):
        try:
            raw = page.extract_text() or ""
        except Exception:
            raw = ""
        text = _normalize_pdf_text(raw)
        if not text.strip():
            pages_without_text.append(page_index)
            continue
        parts = [p.strip() for p in text.split("\n\n") if p.strip()]
        for part_index, part in enumerate(parts, start=1):
            ordinal += 1
            locator_key = f"page:{page_index}#part:{part_index}"
            blocks.append({
                "block_id": _make_block_id(
                    document_id, sha256, "paragraph", locator_key, PDF_PARSER_VERSION
                ),
                "ordinal": ordinal,
                "kind": "paragraph",
                "text": part,
                "locator": {"type": "pdf_page", "page": page_index, "part": part_index},
                "cells": None,
            })

    total_pages = len(reader.pages)
    if not blocks:
        return [], [_error(
            DOCUMENT_PARSE_ERROR, document_id, "local_path",
            "全篇未提取到文本（可能为扫描件或图片型 PDF），需要 OCR/人工核对；不进索引、不算解析成功",
        )], []
    if pages_without_text:
        pages_str = "、".join(str(p) for p in pages_without_text)
        return [], [_error(
            DOCUMENT_PARSE_ERROR, document_id, "local_path",
            f"部分提取：共 {total_pages} 页中第 {pages_str} 页未提取到文本（可能需要 OCR/人工核对，空白页不等同于扫描件）；"
            "为控制范围整篇暂不入索引，请补充资料后重新登记",
        )], []
    return blocks, [], []


def parse_manifest(manifest_path: Path, corpus_root: Path) -> dict:
    """Validate then parse all documents in a manifest."""
    validation = validate_manifest(manifest_path, corpus_root)
    if not validation["valid"]:
        return {
            "schema_version": "1.0",
            "valid": False,
            "document_count": validation["document_count"],
            "documents": [],
            "errors": validation["errors"],
        }

    from campus_assistant.corpus.models import Manifest

    try:
        raw = manifest_path.read_text(encoding="utf-8")
        data: Any = json.loads(raw)
        manifest = Manifest.model_validate(data)
    except Exception:
        return {
            "schema_version": "1.0",
            "valid": False,
            "document_count": 0,
            "documents": [],
            "errors": [_error(MANIFEST_READ_ERROR, None, "manifest", "读取清单失败")],
        }

    doc_results: list[dict] = []
    all_errors: list[dict] = []

    for doc in manifest.documents:
        doc_id = doc.document_id
        ext = Path(doc.local_path).suffix.lower()

        if ext not in SUPPORTED_EXTENSIONS:
            all_errors.append(
                _error(UNSUPPORTED_FORMAT, doc_id, "local_path", f"不支持的文件格式: {ext}")
            )
            continue

        resolved = _resolve_path(corpus_root, doc.local_path)
        if resolved is None or not resolved.is_file():
            all_errors.append(
                _error(DOCUMENT_PARSE_ERROR, doc_id, "local_path", "文件不可访问")
            )
            continue

        from campus_assistant.corpus.validate import _sha256_file

        actual_hash, _ = _sha256_file(resolved)
        if actual_hash != doc.sha256:
            all_errors.append(
                _error(DOCUMENT_PARSE_ERROR, doc_id, "sha256", "文件哈希不符，跳过解析")
            )
            continue

        if ext == ".txt":
            blocks, parse_errors = parse_txt(resolved, doc_id, doc.sha256)
            doc_warnings: list[str] = []
            doc_parser_version = PARSER_VERSION
        elif ext == ".docx":
            blocks, parse_errors, doc_warnings = parse_docx(resolved, doc_id, doc.sha256)
            doc_parser_version = PARSER_VERSION
        elif ext == ".pdf":
            blocks, parse_errors, doc_warnings = parse_pdf(resolved, doc_id, doc.sha256)
            doc_parser_version = PDF_PARSER_VERSION
        else:
            continue

        if parse_errors:
            all_errors.extend(parse_errors)
            continue

        metadata = {
            "document_id": doc.document_id,
            "policy_group_id": doc.policy_group_id,
            "title": doc.title,
            "source_url": doc.source_url,
            "source_site": doc.source_site,
            "local_path": doc.local_path,
            "sha256": doc.sha256,
            "version_label": doc.version_label,
            "published_at": doc.published_at,
            "effective_from": doc.effective_from,
            "effective_to": doc.effective_to,
            "scope": {
                "campus": doc.scope.campus,
                "student_type": doc.scope.student_type,
                "admission_years": doc.scope.admission_years,
            },
            "is_synthetic": doc.is_synthetic,
        }

        result_entry: dict[str, Any] = {
            "document_id": doc_id,
            "metadata": metadata,
            "parser_version": doc_parser_version,
            "blocks": blocks,
            "warnings": doc_warnings,
        }

        doc_results.append(result_entry)

    return {
        "schema_version": "1.0",
        "valid": len(all_errors) == 0,
        "document_count": validation["document_count"],
        "documents": doc_results,
        "errors": all_errors,
    }


def _ensure_utf8_stdout() -> None:
    """R3: Ensure stdout outputs UTF-8 regardless of system locale."""
    if sys.stdout.encoding and sys.stdout.encoding.lower().replace("-", "") == "utf8":
        return
    try:
        sys.stdout = io.TextIOWrapper(
            sys.stdout.buffer, encoding="utf-8", line_buffering=True
        )
    except (AttributeError, OSError):
        pass  # fallback: print with ensure_ascii=True would be the safety net


def main(argv: list[str] | None = None) -> int:
    """CLI entry point. Returns exit code."""
    _ensure_utf8_stdout()

    arg_parser = argparse.ArgumentParser(description="解析语料登记清单中的文档正文")
    arg_parser.add_argument("--manifest", required=True, help="manifest.json 路径")
    arg_parser.add_argument("--corpus-root", required=True, help="语料文件根目录")
    args = arg_parser.parse_args(argv)

    manifest_path = Path(args.manifest)

    try:
        corpus_root = Path(args.corpus_root).resolve()
    except OSError:
        result = {
            "schema_version": "1.0",
            "valid": False,
            "document_count": 0,
            "documents": [],
            "errors": [_error(MANIFEST_READ_ERROR, None, "corpus-root", "无法解析语料根目录")],
        }
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 1

    result = parse_manifest(manifest_path, corpus_root)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["valid"] else 1


if __name__ == "__main__":
    sys.exit(main())
