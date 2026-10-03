"""Astra cross-stage audit. Writes synthetic fixtures only under backend/.cache."""
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import uuid
import zipfile
from importlib.metadata import version

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from fastapi.testclient import TestClient
from campus_assistant.corpus.models import Manifest
from campus_assistant.corpus.parse import parse_manifest
from campus_assistant.corpus.chunk import chunk_manifest
from campus_assistant.corpus.validate import validate_manifest
from campus_assistant.main import create_app
from campus_assistant.retrieval.bm25 import BM25Index

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parents[3]
BACK = ROOT / "step2_code/backend"
TEMP = BACK / ".cache" / ("astra-r3-probes-" + uuid.uuid4().hex)
TEMP.mkdir(parents=True)
example = BACK / "examples/parse-corpus"
seed = json.loads((example / "manifest.json").read_text(encoding="utf-8"))["documents"][0]
out = {}

def entry(path, identity):
    d = dict(seed)
    d.update(document_id=identity, local_path=path.name,
             sha256=hashlib.sha256(path.read_bytes()).hexdigest())
    return d

def manifest(name, entries):
    p = TEMP / name
    p.write_text(json.dumps({"schema_version": "1.0", "documents": entries}, ensure_ascii=False), encoding="utf-8")
    return p

# Healthy chain and source correspondence.
mp = example / "manifest.json"
parsed = parse_manifest(mp, example)
chunked = chunk_manifest(mp, example)
blocks = {(d["document_id"], b["block_id"]): b for d in parsed["documents"] for b in d["blocks"]}
assert all(c["text"] == blocks[(c["document_id"], c["block_id"])]["text"][c["char_start"]:c["char_end"]] for c in chunked["chunks"])
index = BM25Index(chunked["chunks"])
with TestClient(create_app(mp, example)) as client:
    result = client.post("/api/v1/search", json={"query": "缓考申请"})
    assert result.status_code == 200
    assert result.json()["hits"] == index.search("缓考申请")
    assert result.json()["request_id"] == result.headers["X-Request-ID"]
    assert client.get("/api/v1/health/live").status_code == 200
    assert client.get("/api/v1/health/ready").status_code == 503
    out["healthy_chain"] = {"documents": len(parsed["documents"]), "chunks": len(chunked["chunks"]), "hits": len(result.json()["hits"]), "source_slices": "pass", "api_equals_bm25": True}

# A genuinely malformed DOCX passes ZIP opening but fails during table iteration.
good = TEMP / "good.txt"
good.write_text("合成测试 缓考申请", encoding="utf-8")
bad = TEMP / "bad.docx"
doc = Document()
table = doc.add_table(rows=1, cols=1)
table.cell(0, 0).text = "合成测试 表格"
table.cell(0, 0)._tc.get_or_add_tcPr().append(OxmlElement("w:gridSpan"))
doc.save(bad)
bad_mp = manifest("bad-manifest.json", [entry(good, "good"), entry(bad, "bad")])
out["malformed_docx_validation"] = validate_manifest(bad_mp, TEMP)
for name, fn in [("parse", lambda: parse_manifest(bad_mp, TEMP)), ("factory", lambda: create_app(bad_mp, TEMP))]:
    try:
        value = fn()
        out["malformed_docx_" + name] = value if isinstance(value, dict) else {"status": value.state.corpus_status}
    except Exception as exc:
        out["malformed_docx_" + name] = {"uncaught": type(exc).__name__, "message": str(exc)}
cli = subprocess.run([sys.executable, "-m", "campus_assistant.corpus.parse", "--manifest", str(bad_mp), "--corpus-root", str(TEMP)], capture_output=True)
out["malformed_docx_cli"] = {"exit": cli.returncode, "stdout_bytes": len(cli.stdout), "traceback": b"Traceback" in cli.stderr}

# The strict ID contract must reject trailing newline, not silently retain it.
newline = entry(good, "good\n")
newline_mp = manifest("newline-manifest.json", [newline])
out["newline_id_validation"] = validate_manifest(newline_mp, TEMP)
with TestClient(create_app(newline_mp, TEMP)) as client:
    hits = client.post("/api/v1/search", json={"query": "缓考"}).json().get("hits", [])
    out["newline_id_http"] = {"detail_http": client.get("/api/v1/corpus/documents/good%0A").status_code, "preview_url": hits[0]["preview_url"] if hits else None}

# Verify installed dependency versions against lock without reinstalling.
mismatches = []
for line in (BACK / "requirements.lock.txt").read_text().splitlines():
    if "==" in line and not line.startswith("#"):
        package, expected = line.split("==")
        actual = version(package)
        if actual != expected:
            mismatches.append({"package": package, "expected": expected, "actual": actual})
out["lock_mismatches"] = mismatches

# Extract locally built wheel; launch isolated import from a non-source cwd.
wheel = next((BACK / ".cache/astra-full-wheel").glob("*.whl"))
unpacked = TEMP / "wheel"
with zipfile.ZipFile(wheel) as z:
    z.extractall(unpacked)
probe = '''import sys,json
sys.path.insert(0,sys.argv[1])
from campus_assistant.main import create_app
import campus_assistant.main as m
from fastapi.testclient import TestClient
with TestClient(create_app()) as c:
 print(json.dumps({'module':m.__file__,'search':c.get('/search').status_code,'preview':c.get('/preview').status_code}))
'''
p = subprocess.run([sys.executable, "-I", "-c", probe, str(unpacked)], cwd=TEMP, capture_output=True)
out["wheel_non_source"] = {"exit": p.returncode, "result": json.loads(p.stdout) if p.returncode == 0 else p.stderr.decode(errors="replace")}
print(json.dumps(out, ensure_ascii=False, indent=2))
