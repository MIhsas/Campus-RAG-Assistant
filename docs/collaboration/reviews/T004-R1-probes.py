"""Independent synthetic CLI probes; run with backend venv from any cwd."""
import copy
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import uuid
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[3]
backend = root / 'step2_code/backend'
scratch = backend / '.cache' / ('astra-t004-probes-' + uuid.uuid4().hex)
scratch.mkdir(parents=True)
template = json.loads((backend/'examples/parse-corpus/manifest.json').read_text(encoding='utf-8'))['documents'][0]
rows = []

def run(name, path, force_utf8=True):
    entry = copy.deepcopy(template)
    entry.update(document_id=name, local_path=path.name, sha256=hashlib.sha256(path.read_bytes()).hexdigest())
    manifest = scratch/(name+'.json')
    manifest.write_text(json.dumps({'schema_version':'1.0','documents':[entry]},ensure_ascii=False),encoding='utf-8')
    env = dict(os.environ)
    for key in ('PYTHONIOENCODING','PYTHONUTF8'):
        env.pop(key, None)
    if force_utf8:
        env['PYTHONIOENCODING']='utf-8'
    proc = subprocess.run([sys.executable,'-m','campus_assistant.corpus.parse','--manifest',str(manifest),'--corpus-root',str(scratch)],cwd=backend,capture_output=True,env=env)
    row={'case':name,'exit':proc.returncode}
    try:
        out=json.loads(proc.stdout.decode('utf-8'))
        row.update(valid=out['valid'],documents=out['documents'],errors=out['errors'])
    except (UnicodeDecodeError,ValueError) as exc:
        row.update(utf8_json=False,decode_error=type(exc).__name__,stderr_last_line=proc.stderr.decode('utf-8',errors='replace').strip().splitlines()[-1:])
    rows.append(row)
    print(json.dumps(row,ensure_ascii=True))

doc=Document(); table=doc.add_table(rows=3,cols=1)
table.cell(0,0).text='first'; table.cell(2,0).text='third'
path=scratch/'empty-row.docx';doc.save(path);run('empty-row',path)
doc=Document();doc.add_paragraph('body');table=doc.add_table(rows=1,cols=2)
table.cell(0,0).merge(table.cell(0,1)).text='merged'
doc.sections[0].header.paragraphs[0].text='header-not-extracted'
path=scratch/'complex.docx';doc.save(path);run('complex',path)
doc=Document();table=doc.add_table(rows=1,cols=1);table.cell(0,0).text='broken-grid'
table._tbl.remove(table._tbl.tblGrid)
path=scratch/'broken-grid.docx';doc.save(path);run('broken-grid',path)
doc=Document();table=doc.add_table(rows=1,cols=1);table.cell(0,0).text='bad-span'
span=OxmlElement('w:gridSpan');span.set(qn('w:val'),'not-an-integer')
table.cell(0,0)._tc.get_or_add_tcPr().append(span)
path=scratch/'bad-span.docx';doc.save(path);run('bad-span',path)
path=scratch/'unicode.txt';path.write_text('合成测试正文\n',encoding='utf-8');run('default-encoding',path,False)
(root/'docs/collaboration/reviews/T004-R1-evidence.json').write_text(json.dumps(rows,ensure_ascii=True,indent=2),encoding='utf-8')
